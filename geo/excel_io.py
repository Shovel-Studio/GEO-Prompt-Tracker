"""Read prompts from and write results into the tracker workbook.

The Log sheet is rebuilt from prompts.txt (one block of platforms x runs per
prompt). Leaderboard and Domains Cited are recomputed from the Log in Python, so
they are always consistent with whatever has been filled so far.
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from . import config
from .engines.base import Source
from .sources import get_domain

# Log column indices (1-based) — must match template.LOG_HEADERS.
C_DATE, C_PROMPT_NUM, C_PROMPT, C_PLATFORM, C_RUN = 1, 2, 3, 4, 5
C_TOP, C_TOP_BRAND, C_RECS, C_SOURCES, C_ANSWER, C_NOTES = 6, 7, 8, 9, 10, 11

_EXCEL_CELL_MAX = 32000
ANALYSIS_FAILED = "analysis failed"
_NO_KEY = "no analysis key"
_REC_RE = re.compile(r"^\s*(\d+)\.\s*(.*?)(?:\s+—\s+(.*))?\s*$")


def ensure_working_copy(prompts: list[str] | None = None) -> Path:
    """Return the timestamped working copy. Resumes the latest output when it is
    for the same prompt list (or has no results yet); a different prompt list
    gets a fresh file so earlier results are never overwritten."""
    dest = config.resolve_output_path()
    if prompts is not None and dest.exists():
        latest = Tracker(dest)
        if latest._existing_prompts() != prompts and any(latest._completed_rows()):
            dest = config.resolve_output_path(new=True)
    if not dest.exists():
        from .template import build_blank_workbook
        build_blank_workbook(dest)
    return dest


def _format_recs(recs: list[dict]) -> str:
    lines = []
    for i, r in enumerate(recs, start=1):
        same = r["brand"].lower() == r["name"].lower()
        lines.append(f"{i}. {r['name']}" if same else f"{i}. {r['name']} — {r['brand']}")
    return "\n".join(lines)


def _parse_recs(cell) -> list[tuple[int, str, str]]:
    """Inverse of _format_recs: [(position, name, brand), ...]."""
    out = []
    for line in str(cell or "").splitlines():
        m = _REC_RE.match(line)
        if m:
            name = m.group(2).strip()
            out.append((int(m.group(1)), name, (m.group(3) or name).strip()))
    return out


class Tracker:
    def __init__(self, path: Path):
        self.path = path
        self.wb = load_workbook(path)
        self.log = self.wb["Log"]
        self._index: dict[tuple[int, str, int], int] = {}

    # ── structure ────────────────────────────────────────────────────────────
    def sync_structure(self, prompts: list[str]) -> None:
        """Make the Log match `prompts`. Resume-safe: if the prompt set is
        unchanged, existing results are preserved; otherwise the Log is rebuilt."""
        if self._existing_prompts() != prompts:
            self._rebuild_log(prompts)
        self._reindex()

    def _existing_prompts(self) -> list[str]:
        seen, out = set(), []
        for r in range(2, self.log.max_row + 1):
            txt = self.log.cell(r, C_PROMPT).value
            if txt is None:
                continue
            txt = str(txt)
            if txt not in seen:
                seen.add(txt)
                out.append(txt)
        return out

    def _rebuild_log(self, prompts: list[str]) -> None:
        for r in range(2, self.log.max_row + 1):  # clear all data rows
            for c in range(1, C_NOTES + 1):
                self.log.cell(r, c).value = None
        row = 2
        for i, text in enumerate(prompts, start=1):
            for platform in config.PLATFORMS:
                for run in range(1, config.RUNS_PER_PROMPT + 1):
                    self.log.cell(row, C_PROMPT_NUM).value = i
                    self.log.cell(row, C_PROMPT).value = text
                    self.log.cell(row, C_PLATFORM).value = platform
                    self.log.cell(row, C_RUN).value = run
                    row += 1

    def _reindex(self) -> None:
        self._index.clear()
        for r in range(2, self.log.max_row + 1):
            pn = self.log.cell(r, C_PROMPT_NUM).value
            plat = self.log.cell(r, C_PLATFORM).value
            run = self.log.cell(r, C_RUN).value
            if pn is None or plat is None or run is None:
                continue
            self._index[(int(pn), str(plat), int(run))] = r

    # ── read / write ───────────────────────────────────────────────────────────
    def prompts(self) -> list[tuple[int, str]]:
        seen, out = set(), []
        for r in range(2, self.log.max_row + 1):
            pn = self.log.cell(r, C_PROMPT_NUM).value
            if pn is None or int(pn) in seen:
                continue
            seen.add(int(pn))
            out.append((int(pn), str(self.log.cell(r, C_PROMPT).value or "")))
        return out

    def is_filled(self, prompt_num: int, platform: str, run: int) -> bool:
        r = self._index.get((prompt_num, platform, run))
        return bool(r and self.log.cell(r, C_DATE).value)

    def write_result(self, prompt_num: int, platform: str, run: int,
                     sources: list[Source], answer: str, recs: list[dict]) -> None:
        r = self._index.get((prompt_num, platform, run))
        if not r:
            return
        self.log.cell(r, C_DATE).value = date.today().isoformat()
        self.set_recommendations(r, recs)
        self.log.cell(r, C_SOURCES).value = "\n".join(s.url for s in sources)
        self.log.cell(r, C_ANSWER).value = answer[:_EXCEL_CELL_MAX]
        self.log.cell(r, C_NOTES).value = None if config.analysis_key() else _NO_KEY

    def needs_analysis(self) -> list[tuple[int, str, str]]:
        """Answered rows whose recommendations were never extracted:
        [(row, prompt, answer), ...]."""
        out = []
        for r in self._completed_rows():
            note = str(self.log.cell(r, C_NOTES).value or "")
            answer = str(self.log.cell(r, C_ANSWER).value or "")
            if answer and (note.startswith(ANALYSIS_FAILED) or note == _NO_KEY):
                out.append((r, str(self.log.cell(r, C_PROMPT).value or ""), answer))
        return out

    def set_recommendations(self, row: int, recs: list[dict]) -> None:
        self.log.cell(row, C_TOP).value = recs[0]["name"] if recs else ""
        self.log.cell(row, C_TOP_BRAND).value = recs[0]["brand"] if recs else ""
        self.log.cell(row, C_RECS).value = _format_recs(recs)
        self.log.cell(row, C_NOTES).value = None

    def note_error(self, prompt_num: int, platform: str, run: int, msg: str) -> None:
        r = self._index.get((prompt_num, platform, run))
        if r:
            self.log.cell(r, C_NOTES).value = msg

    # ── aggregation ────────────────────────────────────────────────────────────
    def _completed_rows(self):
        for r in range(2, self.log.max_row + 1):
            if self.log.cell(r, C_DATE).value and self.log.cell(r, C_PROMPT_NUM).value is not None:
                yield r

    def leaderboard(self) -> list[dict]:
        """Per prompt, every brand the AIs recommended, most-recommended first."""
        answers: Counter[int] = Counter()
        stats: dict[tuple[int, str], dict] = {}
        prompt_text: dict[int, str] = {}
        for r in self._completed_rows():
            pn = int(self.log.cell(r, C_PROMPT_NUM).value)
            platform = str(self.log.cell(r, C_PLATFORM).value)
            prompt_text[pn] = str(self.log.cell(r, C_PROMPT).value or "")
            answers[pn] += 1
            first_pos: dict[str, int] = {}
            display: dict[str, str] = {}
            items: dict[str, list[str]] = defaultdict(list)
            for pos, name, brand in _parse_recs(self.log.cell(r, C_RECS).value):
                key = brand.lower()
                first_pos.setdefault(key, pos)
                display.setdefault(key, brand)
                items[key].append(name)
            for key, pos in first_pos.items():
                s = stats.setdefault((pn, key), {
                    "brand": display[key], "mentions": 0, "firsts": 0, "positions": [],
                    "platforms": Counter(), "items": Counter()})
                s["mentions"] += 1
                s["firsts"] += pos == 1
                s["positions"].append(pos)
                s["platforms"][platform] += 1
                s["items"].update(n for n in items[key] if n.lower() != key)

        rows = []
        for pn in sorted(prompt_text):
            ranked = sorted(
                (s for (p, _), s in stats.items() if p == pn),
                key=lambda s: (-s["mentions"], -s["firsts"],
                               sum(s["positions"]) / len(s["positions"])))
            for rank, s in enumerate(ranked, start=1):
                rows.append({
                    "prompt_num": pn, "prompt": prompt_text[pn], "rank": rank,
                    "brand": s["brand"], "share": s["mentions"] / answers[pn],
                    "mentions": s["mentions"], "answers": answers[pn], "firsts": s["firsts"],
                    "avg_position": round(sum(s["positions"]) / len(s["positions"]), 1),
                    "platforms": {p: s["platforms"][p] for p in config.PLATFORMS},
                    "items": [n for n, _ in s["items"].most_common()],
                })
        return rows

    def domains(self) -> list[dict]:
        counts: Counter[str] = Counter()
        prompts_for: dict[str, set[int]] = defaultdict(set)
        platforms_for: dict[str, set[str]] = defaultdict(set)
        for r in self._completed_rows():
            pn = int(self.log.cell(r, C_PROMPT_NUM).value)
            platform = str(self.log.cell(r, C_PLATFORM).value)
            for url in str(self.log.cell(r, C_SOURCES).value or "").splitlines():
                dom = get_domain(url)
                if not dom:
                    continue
                counts[dom] += 1
                prompts_for[dom].add(pn)
                platforms_for[dom].add(platform)
        return [{"domain": d, "count": n, "prompts": sorted(prompts_for[d]),
                 "platforms": [p for p in config.PLATFORMS if p in platforms_for[d]]}
                for d, n in counts.most_common()]

    def rebuild_summary_sheets(self) -> None:
        lb = self.wb["Leaderboard"]
        lb.delete_rows(2, lb.max_row)
        for i, row in enumerate(self.leaderboard(), start=2):
            values = [row["prompt_num"], row["prompt"], row["rank"], row["brand"], row["share"],
                      f"{row['mentions']} of {row['answers']}", row["firsts"], row["avg_position"],
                      *row["platforms"].values(), ", ".join(row["items"])]
            for c, v in enumerate(values, start=1):
                lb.cell(i, c).value = v
            lb.cell(i, 5).number_format = "0%"

        d = self.wb["Domains Cited"]
        d.delete_rows(2, d.max_row)
        for i, row in enumerate(self.domains(), start=2):
            d.cell(i, 1).value = row["domain"]
            d.cell(i, 2).value = row["count"]
            d.cell(i, 3).value = ", ".join(str(p) for p in row["prompts"])
            d.cell(i, 4).value = ", ".join(row["platforms"])

    def save(self) -> None:
        self.wb.save(self.path)
