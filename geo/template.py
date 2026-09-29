"""Generate the tracker workbook from scratch so the tool is self-contained.

Three sheets: Log (one row per prompt x platform x run), Leaderboard (which
brands the AIs recommend, per prompt) and Domains Cited (which sources they
cite). The Log is populated from prompts.txt by excel_io.Tracker; Leaderboard
and Domains Cited are rebuilt from the Log after every run.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

LOG_HEADERS = [
    "Date", "Prompt #", "Prompt", "Platform", "Run #",
    "Top Pick", "Top Pick Brand", "Recommendations (ranked, one per line)",
    "Sources Cited (one per line)", "Answer", "Notes",
]

LEADERBOARD_HEADERS = [
    "Prompt #", "Prompt", "Rank", "Brand", "Share of Answers", "Answers Mentioning",
    "Times #1", "Avg Position", "ChatGPT", "Perplexity", "Gemini", "Items Named",
]

DOMAIN_HEADERS = ["Domain", "Times Cited", "Prompts It Appeared On", "Platforms"]


def _header(ws, headers: list[str], row: int = 1) -> None:
    for c, h in enumerate(headers, start=1):
        ws.cell(row, c).value = h
        ws.cell(row, c).font = Font(bold=True)


def build_blank_workbook(path: Path) -> Path:
    wb = Workbook()

    log = wb.active
    log.title = "Log"
    _header(log, LOG_HEADERS)
    log.column_dimensions["C"].width = 40
    log.column_dimensions["F"].width = 28
    log.column_dimensions["H"].width = 40
    log.column_dimensions["I"].width = 40

    lb = wb.create_sheet("Leaderboard")
    _header(lb, LEADERBOARD_HEADERS)
    lb.column_dimensions["B"].width = 36
    lb.column_dimensions["D"].width = 24
    lb.column_dimensions["L"].width = 50

    d = wb.create_sheet("Domains Cited")
    _header(d, DOMAIN_HEADERS)
    d.column_dimensions["A"].width = 34

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
